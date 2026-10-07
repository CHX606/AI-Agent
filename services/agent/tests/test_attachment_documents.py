import base64
from io import BytesIO
from zipfile import ZipFile

import pytest
from bit_agent.attachments import attachment_metadata, validate_attachments, validate_upload_limits
from bit_agent.images import user_message
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def attachment(name, data, mime="application/octet-stream"):
    return {
        "name": name,
        "mime_type": mime,
        "data_url": f"data:{mime};base64,{base64.b64encode(data).decode()}",
    }


def document_bytes(kind):
    stream = BytesIO()
    if kind == "docx":
        document = Document()
        document.add_paragraph("ATTACH-WORD-42 中文")
        document.add_table(rows=1, cols=2).cell(0, 1).text = "table-value"
        document.save(stream)
    elif kind == "xlsx":
        workbook = Workbook()
        workbook.active.title = "测试表"
        workbook.active.append(["ATTACH-SHEET-42", 42, "=B1+1"])
        workbook.save(stream)
        workbook.close()
    else:
        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        contents = DecodedStreamObject()
        contents.set_data(b"BT /F1 12 Tf 10 100 Td (ATTACH-PDF-42) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(contents)
        writer.write(stream)
    return stream.getvalue()


@pytest.mark.parametrize(
    "kind,expected",
    [("docx", "ATTACH-WORD-42 中文"), ("xlsx", "ATTACH-SHEET-42"), ("pdf", "ATTACH-PDF-42")],
)
def test_actual_document_content_reaches_model_without_base64(kind, expected):
    file = attachment(f"资料.{kind}", document_bytes(kind))
    validated = validate_attachments([file])
    blocks = user_message("请分析", attachments=validated)["content"]
    assert blocks[0] == {"type": "input_text", "text": "请分析"}
    assert expected in blocks[1]["text"]
    assert file["data_url"] not in str(blocks)
    assert attachment_metadata(validated)[0]["size"] == len(document_bytes(kind))
    if kind == "docx":
        assert "table-value" in blocks[1]["text"]
    if kind == "xlsx":
        assert "测试表" in blocks[1]["text"] and "=B1+1" in blocks[1]["text"]


@pytest.mark.parametrize(
    "data",
    ["文本代码 = 42".encode(), "中文旧编码".encode("gb18030"), "UTF16文本".encode("utf-16"), b""],
)
def test_text_encodings_and_empty_file(data):
    files = validate_attachments([attachment("note.txt", data)])
    message = user_message("", attachments=files)
    assert message["content"][0]["type"] == "input_text"
    assert "note.txt" in message["content"][0]["text"]
    assert "base64," not in str(message)


@pytest.mark.parametrize("name", ["../secret.txt", "C:secret.txt", "..", "evil\n.txt"])
def test_attachment_names_cannot_be_paths(name):
    with pytest.raises(ValueError):
        validate_attachments([attachment(name, b"safe")])


@pytest.mark.parametrize(
    "name,data",
    [
        ("bad.pdf", b"not a PDF"),
        ("bad.docx", b"broken"),
        ("bad.xlsx", b"broken"),
        ("program.exe", b"MZ\x00binary"),
    ],
)
def test_invalid_or_binary_documents_are_reported(name, data):
    with pytest.raises(ValueError):
        validate_attachments([attachment(name, data)])


def test_combined_limits_and_oversized_text(monkeypatch):
    files = [attachment("note.txt", b"a")]
    with pytest.raises(ValueError, match="合计最多"):
        validate_upload_limits(files * 5, files)
    monkeypatch.setattr("bit_agent.attachments.validation.MAX_TOTAL_BYTES", 1)
    with pytest.raises(ValueError, match="总大小"):
        validate_upload_limits(files, files)
    monkeypatch.setattr("bit_agent.attachments.validation.MAX_TOTAL_BYTES", 20 * 1024 * 1024)
    with pytest.raises(ValueError, match="正文超过"):
        validate_attachments([attachment("large.txt", b"a" * 200_001)])


def test_corrupt_base64_and_mime_mismatch():
    file = attachment("note.txt", b"a")
    with pytest.raises(ValueError, match="规范"):
        validate_attachments([{**file, "data_url": "data:application/octet-stream;base64,YR=="}])
    with pytest.raises(ValueError, match="不一致"):
        validate_attachments([{**file, "mime_type": "text/plain"}])


@pytest.mark.parametrize("suffix", ["docx", "xlsx"])
def test_office_zip_without_required_document_parts_is_reported(suffix):
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr("readme.txt", "this is a ZIP but not an Office document")
    with pytest.raises(ValueError, match="无法读取附件"):
        validate_attachments([attachment(f"broken.{suffix}", stream.getvalue())])
