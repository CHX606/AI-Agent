"""Extract document text with established parsers; never execute uploaded files."""

from collections.abc import Iterator
from io import BytesIO
from pathlib import PurePath
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_TEXT_CHARS = 200_000


def _check_archive(data: bytes) -> None:
    with ZipFile(BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 32 * 1024 * 1024:
            raise ValueError("文档解压后过大，请拆分后上传")


def _pdf(data: bytes) -> Iterator[str]:
    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        raise ValueError("暂不支持加密 PDF，请解密后上传")
    found = False
    for index, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            found = True
            yield f"第 {index} 页\n{text}"
    if not found:
        raise ValueError("PDF 没有可提取的文字；扫描件请上传页面截图")


def _word(data: bytes) -> Iterator[str]:
    _check_archive(data)
    document = Document(BytesIO(data))
    for block in document.iter_inner_content():
        if hasattr(block, "text"):
            yield block.text
        else:
            for row in block.rows:
                yield "\t".join(cell.text for cell in row.cells)


def _spreadsheet(data: bytes) -> Iterator[str]:
    _check_archive(data)
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    try:
        for sheet in workbook:
            if (sheet.max_row or 0) * (sheet.max_column or 0) > 200_000:
                raise ValueError("表格单元格过多，请拆分后上传")
            yield f"工作表：{sheet.title}"
            for row in sheet.iter_rows(values_only=True):
                yield "\t".join("" if value is None else str(value) for value in row)
    finally:
        workbook.close()


def _plain(data: bytes) -> Iterator[str]:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = data.decode("utf-16")
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("gb18030")
    if any(ord(char) < 32 and char not in "\t\r\n\f" for char in text):
        raise ValueError("附件不是可读文本；支持文本、代码、PDF、DOCX 和 XLSX")
    yield text


def extract_text(name: str, data: bytes) -> str:
    suffix = PurePath(name).suffix.lower()
    parser = {".pdf": _pdf, ".docx": _word, ".xlsx": _spreadsheet}.get(suffix, _plain)
    parts, length = [], 0
    try:
        for part in parser(data):
            length += len(part) + 1
            if length > MAX_TEXT_CHARS:
                raise ValueError("附件正文超过 200000 字符，请拆分后上传")
            parts.append(part)
    except (
        UnicodeDecodeError,
        BadZipFile,
        PdfReadError,
        PackageNotFoundError,
        InvalidFileException,
        KeyError,
        ParseError,
        XMLSyntaxError,
    ) as exc:
        raise ValueError(f"无法读取附件 {name}；支持文本、代码、PDF、DOCX 和 XLSX") from exc
    return "\n".join(parts)
